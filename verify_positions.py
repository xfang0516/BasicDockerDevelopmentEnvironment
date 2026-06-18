# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r'd:\文件\unibit\unibit用户操作使用手册\unibit用户操作使用手册.pptx'
prs = Presentation(pptx_path)

REF_TITLE = (2567305, 228600, 4923155, 534035)
REF_CONTENT = (838200, 838200, 9121140, None)
REF_PIC = (609600, 2286000, 8528050, 4492625)

def match(pos, ref, tol=50000):
    return all(abs(pos[i] - ref[i]) <= tol for i in range(min(len(pos), len(ref))))

for idx in range(len(prs.slides)):
    slide = prs.slides[idx]
    title_pos = content_pos = pic_pos = None
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = ''.join(r.text for p in shape.text_frame.paragraphs for r in p.runs)
            if shape.top < 500000:
                title_pos = (shape.left, shape.top, shape.width, shape.height)
            else:
                content_pos = (shape.left, shape.top, shape.width, shape.height)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pic_pos = (shape.left, shape.top, shape.width, shape.height)

    issues = []
    if title_pos and not match(title_pos, REF_TITLE):
        issues.append(f'title pos mismatch: {title_pos}')
    if content_pos and not match(content_pos[:3] + (0,), REF_CONTENT[:3] + (0,)):
        issues.append(f'content pos mismatch: {content_pos}')
    if pic_pos and pic_pos[1] < 2000000:
        issues.append(f'pic too high: {pic_pos}')
    
    status = 'OK' if not issues else '; '.join(issues)
    print(f'Slide {idx+1} ({slide.slide_layout.name}): {status}')
