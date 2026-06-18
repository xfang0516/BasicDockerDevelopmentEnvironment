# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Emu, Pt

pptx_path = r'd:\文件\unibit\unibit用户操作使用手册\unibit用户操作使用手册.pptx'
prs = Presentation(pptx_path)

def emu_to_pt(emu):
    return round(emu / 12700, 1) if emu else None

# Detailed slide 32 analysis
slide = prs.slides[31]
print('=== Slide 32 Detailed ===')
for j, shape in enumerate(slide.shapes):
    print(f'Shape {j}: name={shape.name}, type={shape.shape_type}')
    print(f'  pos: L={shape.left} T={shape.top} W={shape.width} H={shape.height}')
    if shape.has_text_frame:
        tf = shape.text_frame
        print(f'  margins: L={tf.margin_left} T={tf.margin_top} R={tf.margin_right} B={tf.margin_bottom}')
        for k, para in enumerate(tf.paragraphs):
            print(f'  para {k}: level={para.level}, align={para.alignment}, line_spacing={para.line_spacing}')
            pf = para.font
            print(f'    para.font: name={pf.name}, size={pf.size}, bold={pf.bold}')
            for r, run in enumerate(para.runs):
                rf = run.font
                print(f'    run {r}: text={run.text!r}')
                print(f'      font: name={rf.name}, size={emu_to_pt(rf.size)}, bold={rf.bold}')
                try:
                    print(f'      color type={rf.color.type}')
                    if rf.color.rgb:
                        print(f'      rgb={rf.color.rgb}')
                except Exception as e:
                    print(f'      color error: {e}')

# Check GROUP shapes on slide 3
print('\n=== Slide 3 GROUP ===')
slide3 = prs.slides[2]
for j, shape in enumerate(slide3.shapes):
    if str(shape.shape_type).startswith('GROUP'):
        print(f'GROUP shape {j}: name={shape.name}')
        print(f'  pos: L={shape.left} T={shape.top} W={shape.width} H={shape.height}')
        for sub in shape.shapes:
            print(f'  sub: name={sub.name}, type={sub.shape_type}')
            if sub.has_text_frame:
                for para in sub.text_frame.paragraphs:
                    for run in para.runs:
                        print(f'    text={run.text!r}, size={emu_to_pt(run.font.size)}')

# Compare layouts
print('\n=== Layouts ===')
for i, layout in enumerate(prs.slide_layouts):
    print(f'Layout {i}: {layout.name}')
    for j, ph in enumerate(layout.placeholders):
        print(f'  ph {j}: idx={ph.placeholder_format.idx}, type={ph.placeholder_format.type}, name={ph.name}')
        if ph.has_text_frame:
            for para in ph.text_frame.paragraphs:
                pf = para.font
                print(f'    font: size={emu_to_pt(pf.size)}, bold={pf.bold}')

# Summary of all slides - layout and title source
print('\n=== All slides summary ===')
for idx in range(len(prs.slides)):
    slide = prs.slides[idx]
    layout = slide.slide_layout.name
    title_text = ''
    content_preview = ''
    has_group = False
    group_text = ''
    for shape in slide.shapes:
        if str(shape.shape_type).startswith('GROUP'):
            has_group = True
            for sub in shape.shapes:
                if sub.has_text_frame:
                    for para in sub.text_frame.paragraphs:
                        for run in para.runs:
                            group_text += run.text
        if shape.has_text_frame and not str(shape.shape_type).startswith('GROUP'):
            text = ''.join(run.text for para in shape.text_frame.paragraphs for run in para.runs)
            if shape.top < 500000:  # near top - likely title
                if len(text) < 30 and not title_text:
                    title_text = text
            else:
                if not content_preview:
                    content_preview = text[:40]
    print(f'Slide {idx+1}: layout={layout}, group_title={group_text!r}, top_title={title_text!r}, has_group={has_group}')
